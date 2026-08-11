#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""สร้างไฟล์ Excel เล่มเดียวรวมทั้ง ๑๔ ตาราง สำหรับส่งให้คณะกรอก

ผลลัพธ์: แบบฟอร์มกรอกข้อมูลงานวิจัย.xlsx
  แผ่นที่ ๑        คำแนะนำการกรอก + สารบัญแผ่นงาน
  แผ่นที่ ๒–๑๕     ๑๔ ตารางตามเอกสารแนบ ๑ (แถว ๑ = ชื่อฟิลด์, แถว ๒ = คำอธิบาย)

ต้องติดตั้ง openpyxl ก่อน:  pip install openpyxl
วิธีใช้:                    python make_workbook.py
"""

from __future__ import annotations

import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("ไม่พบไลบรารี openpyxl — กรุณาติดตั้งก่อนด้วยคำสั่ง:  pip install openpyxl")

from table_defs import FILL_RULES, SHEET_ORDER_NOTE, TABLES

OUT_PATH = Path(__file__).parent / "แบบฟอร์มกรอกข้อมูลงานวิจัย.xlsx"

FONT = "Tahoma"          # ฟอนต์ที่มีอยู่ทุกเครื่องและแสดงภาษาไทยได้ครบ
HEADER_FILL = "D9E2F3"   # ฟ้าอ่อน — แถวชื่อฟิลด์
DESC_FILL = "F2F2F2"     # เทาอ่อน — แถวคำอธิบาย
KEY_FILL = "FFF2CC"      # เหลืองอ่อน — คอลัมน์ระดับ M (จำเป็น)
DATA_ROWS = 200           # จำนวนแถวที่ตั้งรูปแบบข้อความไว้ล่วงหน้า

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _is_required(desc: str) -> bool:
    return "[M" in desc


def _col_width(field: str, desc: str) -> float:
    """กว้างพอให้เห็นชื่อฟิลด์ครบ แต่ไม่กว้างจนเลื่อนจอ"""
    return max(14.0, min(34.0, len(field) + 4, max(len(field) + 4, len(desc) / 3)))


def _build_guide_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("คำแนะนำ")
    ws.sheet_properties.tabColor = "1F4E79"
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 95

    ws["B1"] = "แบบฟอร์มกรอกข้อมูลผลงานวิจัยและผลงานทางวิชาการ"
    ws["B1"].font = Font(name=FONT, size=14, bold=True, color="1F4E79")
    ws["B2"] = "ระบบ HRCP-KKU-Academic · วิทยาลัยการคอมพิวเตอร์ มหาวิทยาลัยขอนแก่น"
    ws["B2"].font = Font(name=FONT, size=10, color="595959")
    ws["B3"] = ("ใช้ประกอบเอกสารแนบ ๑ (พจนานุกรมข้อมูล) และเอกสารแนบ ๒ "
                "(ข้อกำหนดการส่งมอบข้อมูล) — กรุณาอ่านข้อกำหนดด้านล่างก่อนกรอก")
    ws["B3"].font = Font(name=FONT, size=10, color="595959")

    row = 5
    ws.cell(row=row, column=2, value="ข้อกำหนดการกรอก").font = Font(
        name=FONT, size=12, bold=True)
    row += 1
    for topic, detail in FILL_RULES:
        c_topic = ws.cell(row=row, column=2, value=topic)
        c_topic.font = Font(name=FONT, size=10, bold=True)
        c_topic.alignment = Alignment(vertical="top", wrap_text=True)
        c_topic.border = BORDER
        c_detail = ws.cell(row=row, column=3, value=detail)
        c_detail.font = Font(name=FONT, size=10)
        c_detail.alignment = Alignment(vertical="top", wrap_text=True)
        c_detail.border = BORDER
        ws.row_dimensions[row].height = 15 * (1 + len(detail) // 95)
        row += 1

    row += 1
    ws.cell(row=row, column=2, value="สารบัญแผ่นงาน").font = Font(
        name=FONT, size=12, bold=True)
    row += 1
    ws.cell(row=row, column=3, value=SHEET_ORDER_NOTE).font = Font(
        name=FONT, size=10, italic=True, color="595959")
    row += 2

    c = ws.cell(row=row, column=2, value="แผ่นงาน")
    c.font = Font(name=FONT, size=10, bold=True)
    c.fill = PatternFill("solid", fgColor=HEADER_FILL)
    c.border = BORDER
    c = ws.cell(row=row, column=3, value="ชื่อไทย · จำนวนคอลัมน์ (ในนั้นเป็นฟิลด์จำเป็น M)")
    c.font = Font(name=FONT, size=10, bold=True)
    c.fill = PatternFill("solid", fgColor=HEADER_FILL)
    c.border = BORDER
    row += 1

    for code, thai, cols in TABLES:
        required = sum(1 for _, d in cols if _is_required(d))
        c_code = ws.cell(row=row, column=2, value=code)
        c_code.font = Font(name=FONT, size=10)
        c_code.border = BORDER
        c_desc = ws.cell(row=row, column=3,
                         value=f"{thai} · {len(cols)} คอลัมน์ ({required} ฟิลด์จำเป็น)")
        c_desc.font = Font(name=FONT, size=10)
        c_desc.border = BORDER
        row += 1

    row += 1
    ws.cell(row=row, column=2, value="หมายเหตุ").font = Font(name=FONT, size=10, bold=True)
    ws.cell(row=row, column=3, value=(
        "คอลัมน์ที่แรเงาสีเหลืองคือฟิลด์ระดับ M (จำเป็น) "
        "หากขาดฟิลด์เหล่านี้ ระบบจะเติมแบบฟอร์ม ก.พ.ว. อัตโนมัติไม่ได้"
    )).font = Font(name=FONT, size=10)

    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"


def _build_table_sheet(wb: Workbook, code: str, thai: str,
                       cols: list[tuple[str, str]]) -> None:
    ws = wb.create_sheet(code[:31])
    ws.sheet_properties.tabColor = "2E75B6" if code[0] in "AB" else "70AD47"

    for idx, (field, desc) in enumerate(cols, start=1):
        letter = get_column_letter(idx)
        required = _is_required(desc)

        c_field = ws.cell(row=1, column=idx, value=field)
        c_field.font = Font(name=FONT, size=10, bold=True, color="1F4E79")
        c_field.fill = PatternFill("solid", fgColor=KEY_FILL if required else HEADER_FILL)
        c_field.alignment = Alignment(horizontal="center", vertical="center",
                                      wrap_text=True)
        c_field.border = BORDER

        c_desc = ws.cell(row=2, column=idx, value=desc)
        c_desc.font = Font(name=FONT, size=9, color="595959")
        c_desc.fill = PatternFill("solid", fgColor=DESC_FILL)
        c_desc.alignment = Alignment(vertical="top", wrap_text=True)
        c_desc.border = BORDER

        ws.column_dimensions[letter].width = _col_width(field, desc)

        # ตั้งชนิดเซลล์เป็นข้อความ กัน Excel ตัดศูนย์หน้ารหัส/แปลง ISSN เป็นวันที่
        for r in range(3, DATA_ROWS + 3):
            cell = ws.cell(row=r, column=idx)
            cell.number_format = "@"
            cell.font = Font(name=FONT, size=10)

    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 42
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"

    # ชื่อไทยของตารางใส่ไว้ในชื่อแท็บไม่ได้ (ยาวเกิน) จึงใส่เป็นคำอธิบายแผ่น
    ws.title = code[:31]
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.oddHeader.center.text = f"{code} — {thai}"
    ws.oddHeader.center.size = 10


def main() -> None:
    wb = Workbook()
    wb.remove(wb.active)

    _build_guide_sheet(wb)
    for code, thai, cols in TABLES:
        _build_table_sheet(wb, code, thai, cols)
        print(f"  {code:28s} {len(cols):2d} คอลัมน์")

    wb.save(OUT_PATH)
    total = sum(len(c) for _, _, c in TABLES)
    unique = len({f for _, _, cols in TABLES for f, _ in cols})
    print(f"\nบันทึกแล้ว: {OUT_PATH}")
    print(f"{len(TABLES)} แผ่นตาราง + แผ่นคำแนะนำ ๑ แผ่น")
    print(f"รวม {total} คอลัมน์ ({unique} ชื่อฟิลด์ที่ไม่ซ้ำ — "
          f"เช่น output_id ปรากฏในหลายตารางเพื่อใช้เชื่อมข้อมูล)")


if __name__ == "__main__":
    main()
