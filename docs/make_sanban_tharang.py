from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

DOC_PATH = Path(r"c:/Projects/RM/docs/HRCP_KKU - สำเนา.docx")
FONT_NAME = "TH Sarabun New"
STYLE_NAME = "TableCaption"

BOUNDARY_HEADINGS = {
    "สารบัญ", "สารบัญ(ต่อ)", "สารบัญตาราง", "สารบัญภาพ",
    "บทคัดย่อ", "ABSTRACT", "Abstract",
    "กิตติกรรมประกาศ", "รายการสัญลักษณ์และค าย่อ", "รายการสัญลักษณ์และคำย่อ",
    "บทที่ 1",
}


def set_run_font(run, size_pt: float, bold: bool | None = None) -> None:
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold


def ensure_table_caption_style(doc: Document):
    try:
        style = doc.styles[STYLE_NAME]
    except KeyError:
        style = doc.styles.add_style(STYLE_NAME, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]

    style.font.name = FONT_NAME
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    style.font.size = Pt(14)
    style.font.bold = False
    pf = style.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.0
    return style


def tag_table_captions(doc: Document) -> int:
    count = 0
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("ตารางที่"):
            paragraph.style = doc.styles[STYLE_NAME]
            for run in paragraph.runs:
                set_run_font(run, 14, False)
            count += 1
    return count


def clear_between(doc: Document, start_idx: int, end_idx: int) -> None:
    for idx in range(end_idx - 1, start_idx, -1):
        p = doc.paragraphs[idx]
        p._element.getparent().remove(p._element)


def insert_field_paragraph(after_paragraph, field_code: str) -> None:
    p = after_paragraph.insert_paragraph_before("")
    heading_el = after_paragraph._element
    p_el = p._element
    heading_el.addnext(p_el)

    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")

    text_run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "คลิกขวาแล้วเลือก Update Field"
    text_run.append(t)

    end_char = OxmlElement("w:fldChar")
    end_char.set(qn("w:fldCharType"), "end")

    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text_run)
    run._r.append(end_char)

    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for r in p.runs:
        set_run_font(r, 14, False)


def rebuild_sanban_tharang(doc: Document) -> bool:
    heading_idx = None
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip() == "สารบัญตาราง":
            heading_idx = idx
            break

    if heading_idx is None:
        print("WARNING: 'สารบัญตาราง' heading not found in document.")
        return False

    end_idx = len(doc.paragraphs)
    for i in range(heading_idx + 1, len(doc.paragraphs)):
        if doc.paragraphs[i].text.strip() in BOUNDARY_HEADINGS:
            end_idx = i
            break

    if end_idx - heading_idx > 1:
        clear_between(doc, heading_idx, end_idx)

    heading_par = doc.paragraphs[heading_idx]
    insert_field_paragraph(heading_par, 'TOC \\h \\z \\t "TableCaption,1"')
    return True


def main() -> None:
    if not DOC_PATH.exists():
        raise FileNotFoundError(DOC_PATH)

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = DOC_PATH.with_name(f"HRCP_KKU - สำเนา.backup-{stamp}.docx")
    shutil.copy2(DOC_PATH, backup)
    print(f"backup   {backup}")

    doc = Document(str(DOC_PATH))

    ensure_table_caption_style(doc)
    tagged = tag_table_captions(doc)
    print(f"tagged   {tagged} table caption paragraphs with style '{STYLE_NAME}'")

    ok = rebuild_sanban_tharang(doc)
    if not ok:
        print("Aborted — no สารบัญตาราง heading found. Add the heading and re-run.")
        return

    doc.save(str(DOC_PATH))
    print(f"saved    {DOC_PATH}")
    print()
    print("Next steps:")
    print("  1. Open the document in Microsoft Word")
    print("  2. Press Ctrl+A (select all), then F9 (update fields)")
    print("  3. สารบัญตาราง will fill in with titles and page numbers")
    print("  4. Save the document")


if __name__ == "__main__":
    main()
