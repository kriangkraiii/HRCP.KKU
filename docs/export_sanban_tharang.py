from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

OUT_PATH = Path(r"c:/Projects/RM/docs/สารบัญตาราง.docx")
FONT_NAME = "TH Sarabun New"

COL1_WIDTH = Cm(13.5)  # ชื่อตาราง
COL2_WIDTH = Cm(1.5)   # เลขหน้า

TABLE_ENTRIES = [
    "ตารางที่ 1  สรุปเปรียบเทียบงานวิจัยที่เกี่ยวข้อง",
    "ตารางที่ 2  เปรียบเทียบฟังก์ชันการทำงานของระบบที่เกี่ยวข้อง",
    "ตารางที่ 3  ขั้นตอนการศึกษาและวิเคราะห์ความต้องการของระบบ",
    "ตารางที่ 4  ขั้นตอนการศึกษาและวิเคราะห์ความต้องการของระบบ",
    "ตารางที่ 5  ขั้นตอนการพัฒนาระบบ",
    "ตารางที่ 6  การทดสอบระบบ 3 ระดับ",
    "ตารางที่ 7  ขั้นตอนการติดตั้งและส่งมอบ",
    "ตารางที่ 8  แผนการดำเนินงานตามรอบการพัฒนา",
    "ตารางที่ 9  เครื่องมือด้าน Hardware ที่ใช้",
    "ตารางที่ 10  ภาษาโปรแกรมและ Framework หลักที่ใช้",
    "ตารางที่ 11  เทคโนโลยีด้าน Frontend ที่ใช้",
    "ตารางที่ 12  ระบบจัดการฐานข้อมูลที่ใช้",
    "ตารางที่ 12  ระบบจัดการฐานข้อมูลที่ใช้ (ต่อ)",
    "ตารางที่ 13  เครื่องมือการสร้างและจัดการเอกสารที่ใช้",
    "ตารางที่ 14  ไลบรารีเสริมที่ใช้",
    "ตารางที่ 15  เครื่องมือพัฒนาและ Build ที่ใช้",
    "ตารางที่ 16  เครื่องมือทดสอบที่ใช้",
    "ตารางที่ 17  ตารางสรุปซอฟต์แวร์ทั้งหมด",
    "ตารางที่ 17  ตารางสรุปซอฟต์แวร์ทั้งหมด (ต่อ)",
    "ตารางที่ 17  ตารางสรุปซอฟต์แวร์ทั้งหมด (ต่อ)",
    "ตารางที่ 18  Technology Stack",
    "ตารางที่ 19  โมดูลจัดการผู้ใช้งาน",
    "ตารางที่ 20  โมดูลคำร้องขอประเมินผลงาน",
    "ตารางที่ 21  โมดูลคำร้องขอกำหนดตำแหน่ง",
    "ตารางที่ 22  โมดูลสร้างเอกสาร",
    "ตารางที่ 23  โมดูลจัดการไฟล์",
    "ตารางที่ 24  โมดูลแจ้งเตือน",
    "ตารางที่ 25  ความต้องการเชิงไม่ฟังก์ชัน",
    "ตารางที่ 26  ผู้ใช้งานระบบ",
    "ตารางที่ 27  การยืนยันตัวตน",
    "ตารางที่ 28  การควบคุมการเข้าถึง",
    "ตารางที่ 29  Security Headers",
    "ตารางที่ 30  ข้อมูลทั่วไปของผู้ตอบแบบสอบถามสำรวจความต้องการ",
    "ตารางที่ 31  ปัญหาและอุปสรรคในกระบวนการยื่นเอกสารแบบเดิม (n = 4)",
    "ตารางที่ 32  ระดับความยุ่งยากในการติดตามสถานะเอกสาร",
    "ตารางที่ 33  ฟังก์ชันที่ผู้ใช้ต้องการในระบบ (n = 4)",
    "ตารางที่ 34  ช่องทางการแจ้งเตือนที่ผู้ใช้ต้องการ (n = 4)",
    "ตารางที่ 35  เหตุการณ์ที่ผู้ใช้ต้องการให้ระบบแจ้งเตือน (n = 4)",
    "ตารางที่ 36  ข้อมูลที่ผู้ใช้ต้องการเห็นบนหน้า Dashboard (n = 4)",
    "ตารางที่ 37  อุปกรณ์ที่ผู้ใช้คาดว่าจะใช้งานระบบ (n = 4)",
    "ตารางที่ 38  ขนาดไฟล์ที่มักพบปัญหาเมื่อส่งผ่านระบบออนไลน์ (n = 4)",
    "ตารางที่ 39  ข้อเสนอแนะด้านความปลอดภัยและการเข้าถึงระบบ",
    "ตารางที่ 40  สรุปความต้องการเชิงฟังก์ชันจากการสำรวจ",
    "ตารางที่ 41  สรุปความต้องการเชิงไม่ฟังก์ชันจากการสำรวจ",
    "ตารางที่ 42  ข้อเสนอแนะและความคาดหวังจากผู้ตอบแบบสอบถาม",
    "ตารางที่ 43  usecase description เข้าสู่ระบบ",
    "ตารางที่ 43  usecase description เข้าสู่ระบบ (ต่อ)",
    "ตารางที่ 44  usecase description เข้าสู่ระบบครั้งแรก + OTP",
    "ตารางที่ 44  usecase description เข้าสู่ระบบครั้งแรก + OTP (ต่อ)",
    "ตารางที่ 45  usecase description ลืมรหัสผ่าน",
    "ตารางที่ 46  usecase description รีเซ็ตรหัสผ่าน",
    "ตารางที่ 47  usecase description ออกจากระบบ",
    "ตารางที่ 48  usecase description ดูรายการคำร้องฝั่งผู้ดูแลระบบ",
    "ตารางที่ 49  usecase description ดูรายละเอียดคำร้องฝั่งผู้ดูแลระบบ",
    "ตารางที่ 50  usecase description อัพเดทสถานะคำร้องฝั่งผู้ดูแลระบบ",
    "ตารางที่ 51  usecase description กรอกเอกสารฝั่งผู้ดูแลระบบ",
    "ตารางที่ 51  usecase description กรอกเอกสารฝั่งผู้ดูแลระบบ (ต่อ)",
    "ตารางที่ 52  usecase description บันทึกแบบร่างเอกสารฝั่งผู้ดูแลระบบ",
    "ตารางที่ 53  usecase description ดาวน์โหลดเอกสารฝั่งผู้ดูแลระบบ",
    "ตารางที่ 54  usecase description ดาวน์โหลดทั้งหมด ZIP ฝั่งผู้ดูแลระบบ",
    "ตารางที่ 55  usecase description อัปโหลดเอกสารเพิ่มเติมฝั่งผู้ดูแลระบบ",
    "ตารางที่ 56  usecase description ลบเอกสารเพิ่มเติมฝั่งผู้ดูแลระบบ",
    "ตารางที่ 57  usecase description ส่งข้อเสนอแนะฝั่งผู้ดูแลระบบ",
    "ตารางที่ 58  usecase description ดูแดชบอร์ดฝั่งผู้ยื่นคำร้อง",
    "ตารางที่ 59  usecase description สร้างคำร้องใหม่ฝั่งผู้ยื่นคำร้อง",
    "ตารางที่ 60  usecase description กรอกเอกสารที่ 0 (บันทึกข้อความ) ฝั่งผู้ยื่นคำร้อง",
    "ตารางที่ 61  usecase description กรอกเอกสารที่ 1 (แบบตรวจสอบเบื้องต้น) ฝั่งผู้ยื่นคำร้อง",
    "ตารางที่ 62  usecase description ส่งคำร้องฝั่งผู้ยื่นคำร้อง",
    "ตารางที่ 63  usecase description ดูสถานะคำร้องฝั่งผู้ยื่นคำร้อง",
    "ตารางที่ 64  usecase description อัปโหลดเอกสารแก้ไขฝั่งผู้ยื่นคำร้อง",
    "ตารางที่ 65  usecase description ดูรายการคำร้องตำแหน่งฝั่งผู้ดูแลระบบ",
    "ตารางที่ 66  usecase description ดูรายละเอียดคำร้องตำแหน่งฝั่งผู้ดูแลระบบ",
    "ตารางที่ 67  usecase description อัพเดทสถานะฝั่งผู้ดูแลระบบ",
    "ตารางที่ 68  usecase description กรอกเอกสารที่ 5 & 8 ฝั่งผู้ดูแลระบบ",
    "ตารางที่ 69  usecase description ดูแดชบอร์ดตำแหน่งฝั่งผู้ยื่นคำร้อง",
    "ตารางที่ 70  usecase description สร้างคำร้องตำแหน่งใหม่ฝั่งผู้ยื่นคำร้อง",
    "ตารางที่ 70  usecase description สร้างคำร้องตำแหน่งใหม่ฝั่งผู้ยื่นคำร้อง (ต่อ)",
    "ตารางที่ 71  usecase description กรอกเอกสารฝั่งผู้ยื่นคำร้อง",
    "ตารางที่ 71  usecase description กรอกเอกสารฝั่งผู้ยื่นคำร้อง (ต่อ)",
    "ตารางที่ 72  usecase description ส่งคำร้องตำแหน่งฝั่งผู้ยื่นคำร้อง",
    "ตารางที่ 73  usecase description ดูสถานะคำร้องตำแหน่งฝั่งผู้ยื่นคำร้อง",
    "ตารางที่ 74  usecase description ดูรายชื่อบุคลากร",
    "ตารางที่ 75  usecase description เพิ่มบุคลากร",
    "ตารางที่ 76  usecase description แก้ไขข้อมูลบุคลากร",
    "ตารางที่ 77  usecase description ลบบุคลากร",
    "ตารางที่ 78  usecase description ค้นหาบุคลากรตามบทบาท",
    "ตารางที่ 79  usecase description ดูรายการไฟล์ทั้งหมด",
    "ตารางที่ 80  usecase description ย้ายไฟล์ไปถังขยะ",
    "ตารางที่ 81  usecase description กู้คืนไฟล์จากถังขยะ",
    "ตารางที่ 82  usecase description ลบไฟล์ถาวร",
    "ตารางที่ 83  usecase description ล้างถังขยะ",
    "ตารางที่ 84  usecase description ดาวน์โหลดไฟล์",
    "ตารางที่ 85  usecase description ดูถังขยะ",
    "ตารางที่ 86  usecase description ดูรายชื่อผู้ใช้/แอดมิน",
    "ตารางที่ 87  usecase description เพิ่มแอดมิน",
    "ตารางที่ 88  usecase description แก้ไขข้อมูลผู้ใช้",
    "ตารางที่ 89  usecase description แก้ไขข้อมูลแอดมิน",
    "ตารางที่ 90  usecase description ลบผู้ใช้",
    "ตารางที่ 91  usecase description ลบแอดมิน",
    "ตารางที่ 92  usecase description เปิด/ปิดบัญชี",
    "ตารางที่ 93  usecase description อัพเดทการแจ้งเตือนอีเมล",
    "ตารางที่ 94  usecase description ดูประวัติการใช้งาน (Activity Logs)",
    "ตารางที่ 95  usecase description ดูหน้าตั้งค่า",
    "ตารางที่ 96  usecase description บันทึกการตั้งค่า",
    "ตารางที่ 96  usecase description บันทึกการตั้งค่า (ต่อ)",
    "ตารางที่ 97  usecase description ดูโปรไฟล์",
    "ตารางที่ 98  usecase description แก้ไขโปรไฟล์",
    "ตารางที่ 99  usecase description เปลี่ยนรหัสผ่าน",
    "ตารางที่ 100  usecase description ตรวจสอบหมดอายุอัตโนมัติ",
    "ตารางที่ 101  โครงสร้างตาราง academic_attachment",
    "ตารางที่ 102  โครงสร้างตาราง academic_document",
    "ตารางที่ 103  โครงสร้างตาราง academic_request",
    "ตารางที่ 104  โครงสร้างตาราง admin_logs",
    "ตารางที่ 105  โครงสร้างตาราง notifications",
    "ตารางที่ 106  โครงสร้างตาราง position_request",
    "ตารางที่ 107  โครงสร้างตาราง position_document",
    "ตารางที่ 108  โครงสร้างตาราง position_status_history",
    "ตารางที่ 109  โครงสร้างตาราง request_status_history",
    "ตารางที่ 110  โครงสร้างตาราง staff_member",
    "ตารางที่ 111  โครงสร้างตาราง user_dtls",
    "ตารางที่ 111  โครงสร้างตาราง user_dtls (ต่อ)",
    "ตารางที่ 112  สถานการณ์ทดสอบที่ 1",
    "ตารางที่ 113  สถานการณ์ทดสอบที่ 2",
    "ตารางที่ 114  สถานการณ์ทดสอบที่ 3",
    "ตารางที่ 115  สถานการณ์ทดสอบที่ 4",
    "ตารางที่ 116  สถานการณ์ทดสอบที่ 5",
    "ตารางที่ 117  สถานการณ์ทดสอบที่ 6",
    "ตารางที่ 118  สถานการณ์ทดสอบที่ 7",
    "ตารางที่ 119  สถานการณ์ทดสอบที่ 8",
    "ตารางที่ 120  สถานการณ์ทดสอบที่ 9",
    "ตารางที่ 121  สถานการณ์ทดสอบที่ 10",
    "ตารางที่ 122  สถานการณ์ทดสอบที่ 11",
    "ตารางที่ 123  สถานการณ์ทดสอบที่ 12",
    "ตารางที่ 124  สถานการณ์ทดสอบที่ 13",
    "ตารางที่ 125  สถานการณ์ทดสอบที่ 14",
    "ตารางที่ 126  สถานการณ์ทดสอบที่ 15",
    "ตารางที่ 127  สถานการณ์ทดสอบที่ 16",
    "ตารางที่ 128  สถานการณ์ทดสอบที่ 17",
    "ตารางที่ 129  สถานการณ์ทดสอบที่ 18",
    "ตารางที่ 130  สถานการณ์ทดสอบที่ 19",
    "ตารางที่ 131  สถานการณ์ทดสอบที่ 20",
    "ตารางที่ 132  สถานการณ์ทดสอบที่ 21",
    "ตารางที่ 133  สถานการณ์ทดสอบที่ 22",
    "ตารางที่ 134  สถานการณ์ทดสอบที่ 23",
    "ตารางที่ 135  สถานการณ์ทดสอบที่ 24",
    "ตารางที่ 136  สถานการณ์ทดสอบที่ 25",
    "ตารางที่ 137  สถานการณ์ทดสอบที่ 26",
    "ตารางที่ 138  สถานการณ์ทดสอบที่ 27",
    "ตารางที่ 139  สถานการณ์ทดสอบที่ 28",
    "ตารางที่ 140  สถานการณ์ทดสอบที่ 29",
    "ตารางที่ 141  สถานการณ์ทดสอบที่ 30",
    "ตารางที่ 142  สถานการณ์ทดสอบที่ 31",
    "ตารางที่ 143  สถานการณ์ทดสอบที่ 32",
    "ตารางที่ 144  สถานการณ์ทดสอบที่ 33",
    "ตารางที่ 145  สถานการณ์ทดสอบที่ 34",
    "ตารางที่ 146  สถานการณ์ทดสอบที่ 35",
    "ตารางที่ 147  สถานการณ์ทดสอบที่ 36",
    "ตารางที่ 148  สถานการณ์ทดสอบที่ 37",
    "ตารางที่ 149  สถานการณ์ทดสอบที่ 38",
    "ตารางที่ 150  สถานการณ์ทดสอบที่ 39",
    "ตารางที่ 151  สถานการณ์ทดสอบที่ 40",
    "ตารางที่ 152  สถานการณ์ทดสอบที่ 41",
    "ตารางที่ 153  สถานการณ์ทดสอบที่ 42",
    "ตารางที่ 154  สถานการณ์ทดสอบที่ 43",
    "ตารางที่ 155  สถานการณ์ทดสอบที่ 44",
    "ตารางที่ 156  สถานการณ์ทดสอบที่ 45",
    "ตารางที่ 157  สถานการณ์ทดสอบที่ 46",
    "ตารางที่ 158  สถานการณ์ทดสอบที่ 47",
    "ตารางที่ 159  สถานการณ์ทดสอบที่ 48",
    "ตารางที่ 160  สถานการณ์ทดสอบที่ 49",
    "ตารางที่ 161  สถานการณ์ทดสอบที่ 50",
    "ตารางที่ 162  สถานการณ์ทดสอบที่ 51",
    "ตารางที่ 163  สถานการณ์ทดสอบที่ 52",
    "ตารางที่ 164  แบบสอบถามความพึงพอใจของผู้ใช้งานระบบ",
    "ตารางที่ 165  แบบสอบถามความพึงพอใจของผู้ใช้งานระบบผู้ยื่นคำขอ",
    "ตารางที่ 166  แบบสอบถามความพึงพอใจของผู้ใช้งานระบบผู้ดูแลระบบ",
    "ตารางที่ 167  มาตรวัดแบบ Likert Scale 5 ระดับ",
    "ตารางที่ 168  สถิติที่ใช้ในการวิเคราะห์ข้อมูล",
]


def set_run_font(run, size_pt: float, bold: bool | None = None) -> None:
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold


def remove_table_borders(table) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblBorders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{name}")
        border.set(qn("w:val"), "none")
        tblBorders.append(border)
    tblPr.append(tblBorders)


def style_cell(cell, size_pt: float, bold: bool = False,
               align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT) -> None:
    for p in cell.paragraphs:
        p.alignment = align
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.0
        for run in p.runs:
            set_run_font(run, size_pt, bold)


def main() -> None:
    doc = Document()

    # Page setup: A4 with margins matching the main document
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Inches(1.5)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)

    # Remove the default empty paragraph
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)

    # --- Heading ---
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(12)
    heading.paragraph_format.line_spacing = 1.0
    run = heading.add_run("สารบัญตาราง")
    set_run_font(run, 18, True)

    # --- Table (2 columns: entry | page number) ---
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    remove_table_borders(table)

    for entry in TABLE_ENTRIES:
        row = table.add_row()
        row.cells[0].width = COL1_WIDTH
        row.cells[1].width = COL2_WIDTH
        row.cells[0].text = entry
        row.cells[1].text = ""  # ช่องว่างสำหรับกรอกเลขหน้า
        style_cell(row.cells[0], 14)
        style_cell(row.cells[1], 14, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.save(str(OUT_PATH))
    print(f"saved  {OUT_PATH}")
    print(f"entries: {len(TABLE_ENTRIES)}")


if __name__ == "__main__":
    main()
