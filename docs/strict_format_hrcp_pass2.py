from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

DOC_PATH = Path(r"c:/Projects/RM/docs/HRCP_KKU.docx")
FONT_NAME = "TH Sarabun New"

THAI_LETTER_RANGE = "\u0E00-\u0E7F"

FRONT_TITLES = {
    "บทคัดย่อ",
    "ABSTRACT",
    "Abstract",
    "กิตติกรรมประกาศ",
    "สารบัญ",
    "สารบัญ(ต่อ)",
    "สารบัญตาราง",
    "สารบัญภาพ",
    "รายการสัญลักษณ์และค าย่อ",
    "รายการสัญลักษณ์และคำย่อ",
    "ประวัติผู้เขียน",
}

TOC_HEADINGS = ["สารบัญ", "สารบัญตาราง", "สารบัญภาพ"]
BOUNDARY_HEADINGS = set(TOC_HEADINGS + ["สารบัญ(ต่อ)", "บทที่ 1", "บทคัดย่อ", "ABSTRACT", "Abstract"])


MAIN_HEADING_RE = re.compile(r"^\d+\.\s+")
SUB_HEADING_RE = re.compile(r"^\d+\.\d+")
CHAPTER_RE = re.compile(r"^บทที่\s*\d+\s*$")
TOC_DOTS_RE = re.compile(r"\.{3,}")

PUNCT_SINGLE_RE = re.compile(rf"([\.,:;]) (?=[A-Za-z{THAI_LETTER_RANGE}])")
PUNCT_MULTI_RE = re.compile(rf"([\.,:;])\s{{2,}}(?=[A-Za-z{THAI_LETTER_RANGE}])")


def set_run_font(run, size_pt: float, bold: bool | None = None) -> None:
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold


def apply_text_level_style(paragraph) -> None:
    text = paragraph.text.strip()
    if not text:
        return

    is_chapter = bool(CHAPTER_RE.match(text))
    is_main_heading = bool(MAIN_HEADING_RE.match(text))
    is_sub_heading = bool(SUB_HEADING_RE.match(text))
    is_caption_figure = text.startswith("ภาพที่")
    is_caption_table = text.startswith("ตารางที่")
    is_front = text in FRONT_TITLES

    if is_front or is_chapter:
        for run in paragraph.runs:
            set_run_font(run, 18, True)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_before = Pt(12)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.0
        return

    if is_main_heading:
        for run in paragraph.runs:
            set_run_font(run, 16, True)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.0
        if paragraph.alignment in (None, WD_ALIGN_PARAGRAPH.THAI_JUSTIFY):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        return

    if is_sub_heading:
        for run in paragraph.runs:
            set_run_font(run, 14, False)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.0
        return

    if is_caption_table:
        for run in paragraph.runs:
            set_run_font(run, 14, False)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.0
        return

    if is_caption_figure:
        for run in paragraph.runs:
            set_run_font(run, 14, False)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.0
        return

    if TOC_DOTS_RE.search(text):
        for run in paragraph.runs:
            set_run_font(run, 14, False)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.line_spacing = 1.0
        return

    for run in paragraph.runs:
        set_run_font(run, 14, None)
    paragraph.paragraph_format.first_line_indent = Cm(1.27)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0

    if paragraph.alignment in (None, WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.THAI_JUSTIFY):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.THAI_JUSTIFY


def normalize_punctuation(paragraph) -> int:
    text = paragraph.text
    if not text:
        return 0

    text_stripped = text.strip()
    if (
        text_stripped in FRONT_TITLES
        or CHAPTER_RE.match(text_stripped)
        or MAIN_HEADING_RE.match(text_stripped)
        or SUB_HEADING_RE.match(text_stripped)
        or text_stripped.startswith("ภาพที่")
        or text_stripped.startswith("ตารางที่")
        or TOC_DOTS_RE.search(text_stripped)
    ):
        return 0

    new_text = PUNCT_MULTI_RE.sub(r"\\1  ", text)
    new_text = PUNCT_SINGLE_RE.sub(r"\\1  ", new_text)
    if new_text == text:
        return 0

    first_bold = paragraph.runs[0].bold if paragraph.runs else None
    first_italic = paragraph.runs[0].italic if paragraph.runs else None

    for run in list(paragraph.runs):
        run._r.getparent().remove(run._r)

    run = paragraph.add_run(new_text)
    set_run_font(run, 14, None)
    run.bold = first_bold
    run.italic = first_italic
    return 1


def clear_between(doc: Document, start_idx: int, end_idx: int) -> None:
    for idx in range(end_idx - 1, start_idx, -1):
        p = doc.paragraphs[idx]
        p._element.getparent().remove(p._element)


def insert_field_paragraph(after_paragraph, field_code: str) -> None:
    p = after_paragraph.insert_paragraph_before("")
    # Move inserted paragraph to after the heading by reordering XML
    heading_el = after_paragraph._element
    p_el = p._element
    heading_el.addnext(p_el)

    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")

    text_run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "คลิกขวาแล้วเลือก Update Field"
    text_run.append(t)

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text_run)
    run._r.append(end)

    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0

    for r in p.runs:
        set_run_font(r, 14, False)


def rebuild_toc_blocks(doc: Document) -> None:
    heading_positions = {text.strip(): idx for idx, p in enumerate(doc.paragraphs) for text in [p.text] if text.strip() in TOC_HEADINGS}

    for heading in TOC_HEADINGS:
        if heading not in heading_positions:
            continue
        start = heading_positions[heading]

        end = len(doc.paragraphs)
        for i in range(start + 1, len(doc.paragraphs)):
            t = doc.paragraphs[i].text.strip()
            if t in BOUNDARY_HEADINGS:
                end = i
                break

        if end - start > 1:
            clear_between(doc, start, end)

        heading_par = doc.paragraphs[start]
        if heading == "สารบัญ":
            field = 'TOC \\o "1-3" \\h \\z \\u'
        elif heading == "สารบัญตาราง":
            field = 'TOC \\h \\z \\c "ตาราง"'
        else:
            field = 'TOC \\h \\z \\c "ภาพ"'
        insert_field_paragraph(heading_par, field)


def add_page_number_to_footer(footer_paragraph) -> None:
    run = footer_paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")

    text = OxmlElement("w:t")
    text.text = "1"
    text_run = OxmlElement("w:r")
    text_run.append(text)

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text_run)
    run._r.append(end)


def normalize_sections(doc: Document) -> None:
    for i, section in enumerate(doc.sections):
        section.top_margin = Inches(1.5)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.footer_distance = Inches(1.0)

        section.different_first_page_header_footer = (i == 0)

        footer = section.footer
        footer.is_linked_to_previous = False
        for p in list(footer.paragraphs):
            p._element.getparent().remove(p._element)

        p = footer.add_paragraph("")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        add_page_number_to_footer(p)
        for r in p.runs:
            set_run_font(r, 14, False)


def normalize_tables(doc: Document) -> None:
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run_font(run, 14, None)
                    paragraph.paragraph_format.line_spacing = 1.0


def main() -> None:
    if not DOC_PATH.exists():
        raise FileNotFoundError(DOC_PATH)

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = DOC_PATH.with_name(f"HRCP_KKU.pass2-backup-{stamp}.docx")
    shutil.copy2(DOC_PATH, backup)

    doc = Document(str(DOC_PATH))

    # Base style normalization
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(14)

    normalize_sections(doc)

    punct_changes = 0
    for paragraph in doc.paragraphs:
        punct_changes += normalize_punctuation(paragraph)
        apply_text_level_style(paragraph)

    normalize_tables(doc)
    rebuild_toc_blocks(doc)

    doc.save(str(DOC_PATH))

    print("formatted", DOC_PATH)
    print("backup", backup)
    print("punctuation_changed_paragraphs", punct_changes)


if __name__ == "__main__":
    main()
