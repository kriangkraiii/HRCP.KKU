import os
import re
import sys

try:
    from docx import Document
    from docx.shared import Inches
except ImportError:
    print("python-docx is not installed. Please install it first.")
    sys.exit(1)

def create_word():
    doc = Document()
    doc.add_heading('UAT Test Cases - HRCP-KKU Academic', 0)

    table = doc.add_table(rows=1, cols=6)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'No.'
    hdr_cells[1].text = 'Module'
    hdr_cells[2].text = 'Sub-Module'
    hdr_cells[3].text = 'Test Case ID'
    hdr_cells[4].text = 'Description'
    hdr_cells[5].text = 'Test Result'

    uat_dir = r"c:\Projects\RM\HRCP-KKU-Academic\src\test\java\com\ecom\uat"
    java_files = [f for f in os.listdir(uat_dir) if f.endswith('.java')]

    idx = 1
    for f in java_files:
        content = open(os.path.join(uat_dir, f), 'r', encoding='utf-8').read()
        
        main_mod_match = re.search(r'@DisplayName\("([^"]+)"\)\s*class', content)
        main_mod = main_mod_match.group(1) if main_mod_match else f
        if "UAT: " in main_mod:
            main_mod = main_mod.replace("UAT: ", "")
        
        # Split by @Nested
        nested_blocks = content.split('@Nested')
        if len(nested_blocks) > 1:
            nested_blocks = nested_blocks[1:]
        else:
            # If no @Nested, just process the whole class
            nested_blocks = [content]
            
        for block in nested_blocks:
            display_names = re.findall(r'@DisplayName\("([^"]+)"\)', block)
            if not display_names:
                continue
                
            sub_mod_full = display_names[0]
            if ': ' in sub_mod_full:
                sub_mod_id, sub_mod_name = sub_mod_full.split(': ', 1)
            else:
                sub_mod_id, sub_mod_name = "", sub_mod_full
                
            test_cases = display_names[1:]
            for test_full in test_cases:
                if ': ' in test_full:
                    t_id, t_desc = test_full.split(': ', 1)
                else:
                    t_id, t_desc = "", test_full
                    
                full_tc_id = f"{sub_mod_id}-{t_id}" if sub_mod_id and t_id else (t_id or sub_mod_id)
                
                row = table.add_row().cells
                row[0].text = str(idx)
                row[1].text = main_mod
                row[2].text = sub_mod_name
                row[3].text = full_tc_id
                row[4].text = t_desc
                row[5].text = 'Pass'
                idx += 1

    out_path = r'c:\Projects\RM\HRCP-KKU-Academic\UAT_Test_Cases.docx'
    doc.save(out_path)
    print(f"Generated {out_path} with {idx-1} test cases.")

if __name__ == '__main__':
    create_word()
