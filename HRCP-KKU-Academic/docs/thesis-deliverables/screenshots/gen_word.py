# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import os

BASE = r"c:/Projects/RM/HRCP-KKU-Academic/docs/screenshots/sequence diagram"
OUT  = r"c:/Projects/RM/HRCP-KKU-Academic/docs/screenshots/Sequence_Diagrams.docx"

doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin   = Cm(3.17)
    section.right_margin  = Cm(3.17)

styles = doc.styles

def ensure_style(name, bold=False, size=None, color=None, sb=None, sa=None, italic=False):
    try:
        st = styles[name]
    except KeyError:
        st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = styles["Normal"]
    pf = st.paragraph_format
    if sb is not None: pf.space_before = Pt(sb)
    if sa is not None: pf.space_after  = Pt(sa)
    rf = st.font
    rf.name = "TH SarabunPSK"
    if bold:   rf.bold   = bold
    if italic: rf.italic = italic
    if size:   rf.size   = Pt(size)
    if color:  rf.color.rgb = RGBColor(*color)
    return st

ensure_style("TH H1",      bold=True,  size=20, color=(31,73,125),  sb=14, sa=6)
ensure_style("TH H2",      bold=True,  size=15, color=(68,114,196), sb=10, sa=4)
ensure_style("TH Caption", bold=True,  size=11, color=(89,89,89),   sb=6,  sa=3, italic=True)
ensure_style("TH Body",    bold=False, size=14, color=(0,0,0),      sb=0,  sa=8)

def h1(text):
    p = doc.add_paragraph(text, style="TH H1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

def h2(text):
    doc.add_paragraph(text, style="TH H2")

def caption(fig_no, text):
    p = doc.add_paragraph(style="TH Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"\u0e23\u0e39\u0e1b\u0e17\u0e35\u0e48 {fig_no}: {text}")

def body(text):
    p = doc.add_paragraph(text, style="TH Body")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

def image(path, width=Inches(5.8)):
    if os.path.exists(path):
        doc.add_picture(path, width=width)
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        doc.add_paragraph(f"[ไม่พบไฟล์: {path}]")

# ── Cover ──
for _ in range(6): doc.add_paragraph("")
h1("Sequence Diagram")
h1("ระบบ HRCP-KKU Academic")
doc.add_paragraph("")
p = doc.add_paragraph("เอกสารนี้แสดง Sequence Diagram ของ Workflow หลักในระบบ", style="TH Body")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph("ประกอบด้วย Academic Evaluation และ Position Request", style="TH Body")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

fig = 1
ACAD = BASE + "/Academic Evaluation"
POS  = BASE + "/Position Request"

# ══════════════════════════════════════════════
# SECTION 1 : ACADEMIC EVALUATION
# ══════════════════════════════════════════════
h1("ส่วนที่ 1: Academic Evaluation (ประเมินผลการสอน)")
body(
    "กระบวนการยื่นคำร้องขอประเมินผลการสอน เริ่มตั้งแต่ผู้ยื่นสร้างคำร้องและกรอกเอกสาร "
    "จนกระทั่ง Admin ดำเนินการครบถ้วนและแจ้งผลการประเมิน โดยแบ่งเป็นเอกสาร 9 ประเภท (Doc 0–8)"
)
doc.add_page_break()

acad = [
    (
        "Doc 0 \u2014 \u0e1a\u0e31\u0e19\u0e17\u0e36\u0e01\u0e02\u0e49\u0e2d\u0e04\u0e27\u0e32\u0e21 (\u0e1c\u0e39\u0e49\u0e22\u0e37\u0e48\u0e19).drawio.png",
        "Doc 0 — บันทึกข้อความ (ผู้ยื่นคำร้อง)",
        "Doc 0 คือเอกสารแรกที่ผู้ยื่นต้องกรอก ได้แก่ บันทึกข้อความขอรับการประเมินผลการสอน "
        "ผู้ยื่นสามารถบันทึกร่าง (Draft) ไว้ก่อนได้ หรือยืนยันเพื่อให้ระบบสร้างไฟล์ DOCX โดยอัตโนมัติ "
        "ข้อมูลที่กรอกในเอกสารนี้จะถูกนำไป auto-fill ในเอกสารชุดถัดไปของ Admin"
    ),
    (
        "Doc 1 \u2014 \u0e41\u0e1a\u0e1a\u0e15\u0e23\u0e27\u0e08\u0e2a\u0e2d\u0e1a\u0e40\u0e1a\u0e37\u0e49\u0e2d\u0e07\u0e15\u0e49\u0e19 (\u0e1c\u0e39\u0e49\u0e22\u0e37\u0e48\u0e19).drawio.png",
        "Doc 1 — แบบตรวจสอบเบื้องต้น (ผู้ยื่นคำร้อง)",
        "Doc 1 คือแบบตรวจสอบเบื้องต้นที่ผู้ยื่นต้องยืนยันความถูกต้องของข้อมูลก่อนส่งคำร้อง "
        "มีกลไก Draft เช่นเดียวกับ Doc 0 เมื่อยืนยันแล้ว ระบบจะ generate ไฟล์ DOCX และบันทึกลงฐานข้อมูล"
    ),
    (
        "Submit Request \u2014 \u0e2a\u0e48\u0e07\u0e04\u0e33\u0e23\u0e49\u0e2d\u0e07.drawio.png",
        "Submit Request — ส่งคำร้องประเมินผลการสอน",
        "เมื่อกรอกเอกสารครบแล้ว ผู้ยื่นกดส่งคำร้อง ระบบจะเปลี่ยนสถานะเป็น RECEIVED "
        "บันทึกวันที่ส่ง และส่งอีเมลแจ้งเตือนไปยัง Admin ทุกคนที่เปิดใช้งานการแจ้งเตือนอีเมล"
    ),
    (
        "Doc 2 \u2014 \u0e04\u0e33\u0e02\u0e2d\u0e41\u0e15\u0e48\u0e07\u0e15\u0e31\u0e49\u0e07\u0e01\u0e23\u0e23\u0e21\u0e01\u0e32\u0e23 (Admin).drawio.png",
        "Doc 2 — คำขอแต่งตั้งกรรมการ (Admin)",
        "Admin กรอกรายชื่อกรรมการที่จะแต่งตั้ง ข้อมูลกรรมการที่ระบุไว้ที่นี่จะถูกนำไป auto-fill "
        "ในเอกสาร Doc 3 และ Doc 4 ต่อไป"
    ),
    (
        "Doc 3 \u2014 \u0e04\u0e33\u0e2a\u0e31\u0e48\u0e07\u0e41\u0e15\u0e48\u0e07\u0e15\u0e31\u0e49\u0e07\u0e01\u0e23\u0e23\u0e21\u0e01\u0e32\u0e23 (Admin).drawio.png",
        "Doc 3 — คำสั่งแต่งตั้งกรรมการ (Admin) → SUB_COMMITTEE_APPOINTED",
        "เมื่อ Admin บันทึก Doc 3 และเลือก sendNotify ระบบจะเปลี่ยนสถานะอัตโนมัติเป็น "
        "SUB_COMMITTEE_APPOINTED พร้อมส่งอีเมลแจ้งผู้ยื่นว่าได้แต่งตั้งกรรมการแล้ว"
    ),
    (
        "Doc 4 \u2014 \u0e2b\u0e19\u0e31\u0e07\u0e2a\u0e37\u0e2d\u0e40\u0e0a\u0e34\u0e0d\u0e01\u0e23\u0e23\u0e21\u0e01\u0e32\u0e23 3 \u0e0a\u0e38\u0e14 (Admin).drawio.png",
        "Doc 4 — หนังสือเชิญกรรมการ 3 ชุด (Admin) → MEETING_SCHEDULED",
        "Doc 4 จะถูก generate เป็น 3 ฉบับแยกกัน (copyNumber 1–3) สำหรับกรรมการแต่ละท่าน "
        "เมื่อบันทึกพร้อม sendNotify สถานะจะเปลี่ยนเป็น MEETING_SCHEDULED "
        "และแจ้งผู้ยื่นพร้อมวันที่และสถานที่ประชุม"
    ),
    (
        "Doc 5 \u2014 \u0e23\u0e32\u0e22\u0e07\u0e32\u0e19\u0e01\u0e32\u0e23\u0e1b\u0e23\u0e30\u0e0a\u0e38\u0e21 + \u0e2a\u0e48\u0e07\u0e02\u0e49\u0e2d\u0e40\u0e2a\u0e19\u0e2d\u0e41\u0e19\u0e30 (Admin).drawio.png",
        "Doc 5 — รายงานการประชุมและการส่งข้อเสนอแนะ (Admin)",
        "Admin กรอกรายงานการประชุม หากมีข้อเสนอแนะให้แก้ไข สามารถกด 'ส่งข้อเสนอแนะ' "
        "ซึ่งจะเปลี่ยนสถานะเป็น COMPLETED_REVISE และส่งอีเมล 2 ฉบับ ได้แก่ "
        "อีเมลแจ้งสถานะและอีเมลรายละเอียดข้อเสนอแนะในกล่องสีส้ม "
        "จากนั้นผู้ยื่นสามารถอัปโหลดเอกสารที่แก้ไขแล้วผ่านหน้า upload-revision"
    ),
    (
        "Doc 6 \u2014 \u0e41\u0e1a\u0e1a\u0e1b\u0e23\u0e30\u0e40\u0e21\u0e34\u0e19\u0e1c\u0e25\u0e01\u0e32\u0e23\u0e2a\u0e2d\u0e19 (Admin).drawio.png",
        "Doc 6 — แบบประเมินผลการสอน (Admin) → COMPLETED_PASS / COMPLETED_FAIL",
        "Admin กรอกคะแนนประเมิน 4 ส่วน ระบบคำนวณคะแนนถ่วงน้ำหนักอัตโนมัติ "
        "(ส่วน 1,4 = 20%, ส่วน 2,3 = 30%) แล้วตัดสินผล: ต่ำกว่า 57 = ไม่ผ่าน, 57–70 = ชำนาญ, "
        "71–85 = ชำนาญพิเศษ, 86+ = เชี่ยวชาญ เมื่อ sendNotify สถานะจะเปลี่ยนเป็น "
        "COMPLETED_PASS หรือ COMPLETED_FAIL พร้อมส่งอีเมลสีเขียว/แดงตามผล"
    ),
    (
        "Doc 7 \u2014 \u0e43\u0e1a\u0e23\u0e31\u0e1a\u0e23\u0e2d\u0e07\u0e1c\u0e25\u0e01\u0e32\u0e23\u0e1b\u0e23\u0e30\u0e0a\u0e38\u0e21 (Admin).drawio.png",
        "Doc 7 — ใบรับรองผลการประชุม (Admin)",
        "ระบบ auto-fill ข้อมูลจาก Doc 6 มายัง Doc 7 พร้อมแปลงเลขลำดับการประชุมเป็นเลขไทย "
        "Admin ตรวจสอบและยืนยันเพื่อสร้างไฟล์ DOCX"
    ),
    (
        "Doc 8 \u2014 \u0e2b\u0e19\u0e31\u0e07\u0e2a\u0e37\u0e2d\u0e41\u0e08\u0e49\u0e07\u0e1c\u0e25\u0e01\u0e32\u0e23\u0e1b\u0e23\u0e30\u0e40\u0e21\u0e34\u0e19 (Admin).drawio.png",
        "Doc 8 — หนังสือแจ้งผลการประเมิน (Admin) → COMPLETED",
        "Doc 8 คือเอกสารสุดท้ายของ workflow โดยระบบ auto-fill จาก Doc 7 และแปลงวันที่เป็นเลขไทย "
        "เมื่อ Admin บันทึกพร้อม sendNotify สถานะคำร้องจะเปลี่ยนเป็น COMPLETED "
        "และส่งอีเมลแจ้งผู้ยื่นว่าดำเนินการเสร็จสิ้น"
    ),
]

h2("1.1 Academic Evaluation — Sequence Diagrams แยกตามเอกสาร")
for fname, cap, desc in acad:
    image(os.path.join(ACAD, fname))
    caption(fig, cap)
    body(desc)
    doc.add_paragraph("")
    fig += 1

doc.add_page_break()

# ══════════════════════════════════════════════
# SECTION 2 : POSITION REQUEST
# ══════════════════════════════════════════════
h1("ส่วนที่ 2: Position Request (ตำแหน่งทางวิชาการ)")
body(
    "กระบวนการยื่นคำร้องขอตำแหน่งทางวิชาการ เริ่มตั้งแต่ผู้ยื่นตรวจสอบสิทธิ์ กรอกเอกสาร 7 ประเภท "
    "ส่งคำร้อง จนถึง Admin ดำเนินการผ่านขั้นตอนคณะกรรมการและส่งเรื่องให้งาน HR "
    "มีกลไก REVISION_REQUESTED ให้ Admin ขอให้ผู้ยื่นแก้ไขเอกสารได้ทุกขั้นตอน"
)
doc.add_page_break()

pos = [
    (
        "\u0e2a\u0e23\u0e49\u0e32\u0e07\u0e04\u0e33\u0e23\u0e49\u0e2d\u0e07 + \u0e15\u0e23\u0e27\u0e08\u0e2a\u0e2d\u0e1a\u0e2a\u0e34\u0e17\u0e18\u0e34\u0e4c.drawio.png",
        "สร้างคำร้อง + ตรวจสอบสิทธิ์ (ผู้ยื่นคำร้อง)",
        "ก่อนสร้างคำร้อง ระบบตรวจสอบ 2 เงื่อนไข ได้แก่ (1) ผู้ยื่นต้องไม่มีคำร้องที่กำลังดำเนินการอยู่ "
        "และ (2) ต้องมีผลการประเมินการสอนที่ผ่านแล้วและ Doc 8 ยังไม่หมดอายุ "
        "เมื่อสร้างสำเร็จ ระบบออกรหัสคำร้องรูปแบบ KKU-POS-{ปีพุทธศักราช}-{id}"
    ),
    (
        "Doc 1 \u2014 \u0e43\u0e1a\u0e2a\u0e21\u0e31\u0e04\u0e23 _ \u0e02\u0e49\u0e2d\u0e21\u0e39\u0e25\u0e2a\u0e48\u0e27\u0e19\u0e15\u0e31\u0e27 (\u0e1c\u0e39\u0e49\u0e22\u0e37\u0e48\u0e19).drawio.png",
        "Doc 1 — ใบสมัคร / ข้อมูลส่วนตัว (ผู้ยื่นคำร้อง)",
        "Doc 1 คือเอกสารข้อมูลส่วนตัวของผู้ยื่น ประกอบด้วย ชื่อ-นามสกุล ตำแหน่ง สังกัด และข้อมูลการศึกษา "
        "ข้อมูลจาก Doc 1 จะถูกนำไป auto-fill ในเอกสารประเภทอื่น ๆ ทั้งหมด "
        "ระบบบันทึก EditLog ทุกครั้งที่มีการบันทึกร่างหรือยืนยัน"
    ),
    (
        "Doc 2 \u2014 \u0e02\u0e49\u0e2d\u0e21\u0e39\u0e25\u0e15\u0e33\u0e41\u0e2b\u0e19\u0e48\u0e07\u0e17\u0e35\u0e48\u0e02\u0e2d (\u0e1c\u0e39\u0e49\u0e22\u0e37\u0e48\u0e19).drawio.png",
        "Doc 2 — ข้อมูลตำแหน่งที่ขอ (ผู้ยื่นคำร้อง)",
        "Doc 2 ระบุตำแหน่งทางวิชาการที่ต้องการขอ สาขาวิชา และวิธีการประเมิน "
        "เมื่อยืนยัน ระบบจะ sync ข้อมูล targetPosition, major, evaluationMethod "
        "กลับไปอัปเดตใน entity PositionRequest เพื่อใช้แสดงผลในรายการ"
    ),
    (
        "Doc 3, 4, 6, 7, 9 \u2014 \u0e40\u0e2d\u0e01\u0e2a\u0e32\u0e23\u0e1b\u0e23\u0e30\u0e01\u0e2d\u0e1a\u0e2d\u0e37\u0e48\u0e19 \u0e46 (\u0e1c\u0e39\u0e49\u0e22\u0e37\u0e48\u0e19).drawio.png",
        "Doc 3, 4, 6, 7, 9 — เอกสารประกอบอื่น ๆ (ผู้ยื่นคำร้อง)",
        "เอกสารประกอบส่วนที่เหลือ ได้แก่ ผลงานทางวิชาการ เอกสารคุณสมบัติ และหลักฐานต่าง ๆ "
        "ทุกเอกสารรองรับการบันทึกร่างและยืนยัน โดยระบบโหลด Doc 1 มา auto-fill "
        "ชื่อ-ตำแหน่งผู้ยื่นให้อัตโนมัติ และบันทึก EditLog ทุกครั้ง"
    ),
    (
        "Submit Request \u2014 \u0e2a\u0e48\u0e07\u0e04\u0e33\u0e23\u0e49\u0e2d\u0e07 position.drawio.png",
        "Submit Request — ส่งคำร้องตำแหน่งทางวิชาการ",
        "ก่อนส่ง ระบบตรวจสอบว่าเอกสารที่จำเป็นทั้ง 7 ประเภท (Doc 1, 2, 3, 4, 6, 7, 9) "
        "ครบถ้วนและยืนยันแล้ว หากเอกสารไม่ครบจะแจ้ง error พร้อมระบุ Doc ที่ขาด "
        "เมื่อส่งสำเร็จสถานะจะเปลี่ยนเป็น DOCUMENT_RECEIVED และส่งอีเมลแจ้ง Admin"
    ),
    (
        "Doc 5 \u2014 \u0e1b\u0e23\u0e30\u0e40\u0e21\u0e34\u0e19\u0e42\u0e14\u0e22\u0e1c\u0e39\u0e49\u0e1a\u0e31\u0e07\u0e04\u0e31\u0e1a\u0e1a\u0e31\u0e0d\u0e0a\u0e32 (Admin).drawio.png",
        "Doc 5 — ประเมินโดยผู้บังคับบัญชา (Admin)",
        "Doc 5 เป็นเอกสาร Admin เท่านั้น ผู้ยื่นไม่สามารถเข้าถึงได้ "
        "Admin กรอกผลการประเมินจากผู้บังคับบัญชา ระบบพยายาม generate DOCX "
        "หากล้มเหลวยังคงบันทึก JSON ลงฐานข้อมูลได้ โดยไม่มี filePath"
    ),
    (
        "Status Updates \u2014 \u0e01\u0e32\u0e23\u0e40\u0e1b\u0e25\u0e35\u0e48\u0e22\u0e19\u0e2a\u0e16\u0e32\u0e19\u0e30 (Admin).drawio.png",
        "Status Updates — การเปลี่ยนสถานะคำร้องตำแหน่ง (Admin)",
        "Admin เปลี่ยนสถานะคำร้องผ่าน Flow: DOCUMENT_RECEIVED → DOCUMENT_VERIFICATION → "
        "SCREENING_COMMITTEE → SCREENING_APPROVED → COLLEGE_COMMITTEE → COLLEGE_APPROVED → SENT_TO_HR "
        "ทุกการเปลี่ยนสถานะจะบันทึก StatusHistory และส่งอีเมลแจ้งผู้ยื่น "
        "โดยสีอีเมลต่างกันตามสถานะ: ส้ม = REVISION_REQUESTED, เขียว = ผ่าน, เทียล = SENT_TO_HR "
        "Admin ยังสามารถขอให้ผู้ยื่นแก้ไขเอกสาร (REVISION_REQUESTED) ได้ทุกขั้นตอน"
    ),
    (
        "Doc 8 \u2014 \u0e2a\u0e23\u0e38\u0e1b\u0e1c\u0e25\u0e41\u0e25\u0e30\u0e1a\u0e31\u0e0d\u0e0a\u0e35\u0e23\u0e32\u0e22\u0e0a\u0e37\u0e48\u0e2d\u0e1c\u0e39\u0e49\u0e21\u0e35\u0e04\u0e38\u0e13\u0e2a\u0e21\u0e1a\u0e31\u0e15\u0e34 (Admin).drawio.png",
        "Doc 8 — สรุปผลและบัญชีรายชื่อผู้มีคุณสมบัติ (Admin)",
        "Doc 8 สรุปผลการพิจารณาของคณะกรรมการ ระบบ auto-fill ข้อมูลจาก Doc 6 "
        "และ merge ข้อมูลใหม่กับข้อมูลเดิมเพื่อป้องกันการสูญหายของ field ที่ไม่ได้แก้ไข "
        "เป็นเอกสาร Admin เท่านั้น"
    ),
    (
        "Download Documents \u2014 \u0e14\u0e32\u0e27\u0e19\u0e4c\u0e42\u0e2b\u0e25\u0e14\u0e40\u0e2d\u0e01\u0e2a\u0e32\u0e23 (Admin).drawio.png",
        "Download Documents — ดาวน์โหลดเอกสาร (Admin)",
        "Admin ดาวน์โหลดเอกสารได้ 2 แบบ: แบบเดี่ยวต่อ docType หรือดาวน์โหลดทั้งหมดเป็น ZIP "
        "ทั้งสองแบบรองรับการ generate เอกสารใหม่อัตโนมัติหากไม่พบไฟล์บน disk "
        "ข้อแตกต่างสำคัญ: Download เดี่ยวจะ save generatedFilePath กลับลงฐานข้อมูล (cache) "
        "แต่ Download ZIP ไม่บันทึก path กลับ"
    ),
    (
        "Attachments \u2014 \u0e08\u0e31\u0e14\u0e01\u0e32\u0e23\u0e44\u0e1f\u0e25\u0e4c\u0e41\u0e19\u0e1a (Admin).drawio.png",
        "Attachments — จัดการไฟล์แนบ (Admin)",
        "Admin อัปโหลด ดาวน์โหลด และลบไฟล์แนบประกอบคำร้อง "
        "รองรับเฉพาะไฟล์ PDF และ DOCX จำกัดสูงสุด 10 ไฟล์ต่อคำร้อง "
        "การลบใช้ Soft Delete โดยตั้ง isDeleted=true แทนการลบจริง "
        "ชื่อไฟล์ที่จัดเก็บมี timestamp นำหน้าเพื่อป้องกันชื่อซ้ำ"
    ),
]

h2("2.1 Position Request — Sequence Diagrams แยกตามเอกสาร")
for fname, cap, desc in pos:
    image(os.path.join(POS, fname))
    caption(fig, cap)
    body(desc)
    doc.add_paragraph("")
    fig += 1

doc.save(OUT)
print(f"Done: {OUT}  ({fig-1} diagrams)")
